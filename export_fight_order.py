"""Write fight_order.json so the website lists this week's bouts in order.

The order comes from UFC Stats: the upcoming event whose fighters best match
DKSalaries.csv, top of the card first (the main event is Fight 1). If UFC Stats
can't be reached or no event matches the slate, the 'DraftKings Fighter Pool'
sheet of the sim workbook (column A = Fighter, column B = FightID) is used.
publish.bat runs this before it stages the contest files.

Usage:  python export_fight_order.py ["Post Contest Sim.xlsm"] [--workbook]
        --workbook   skip UFC Stats and use the workbook's FightIDs
"""
import csv
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
SHEET = "DraftKings Fighter Pool"


def norm(name):
    s = unicodedata.normalize("NFD", str(name or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\s+(jr|sr|ii|iii|iv)\.?$", "", s.strip())
    return re.sub(r"[^a-z]", "", s)


def slate_names():
    path = HERE / "DKSalaries.csv"
    if not path.exists():
        return set()
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {norm(r.get("Name")) for r in csv.DictReader(f) if r.get("Name")}


def from_ufcstats(slate):
    """The best-matching upcoming event's bouts, top of the card first, or None."""
    try:
        sys.path.insert(0, str(HERE))
        import scrape_ufc_stats as us   # reuses its session and browser-check solver
        from bs4 import BeautifulSoup
    except Exception as e:
        print(f"export_fight_order: UFC Stats unavailable ({e}); using the workbook.")
        return None
    html = us.fetch(us.BASE_URL + "/statistics/events/upcoming")
    if not html:
        print("export_fight_order: could not load UFC Stats upcoming events; using the workbook.")
        return None
    links = []
    for a in BeautifulSoup(html, "html.parser").select("a.b-link"):
        href = a.get("href", "")
        if "event-details" in href and href not in [l[1] for l in links]:
            links.append((a.get_text(strip=True), href))
    best = None
    for title, href in links[:4]:
        page = us.fetch(href)
        if not page:
            continue
        bouts = []
        for tr in BeautifulSoup(page, "html.parser").select("tr.b-fight-details__table-row"):
            names = [a.get_text(strip=True) for a in tr.select("td:nth-of-type(2) a")]
            if len(names) >= 2:
                bouts.append(names[:2])
        hits = sum(1 for b in bouts for n in b if norm(n) in slate)
        if bouts and (best is None or hits > best[0]):
            best = (hits, title, bouts)
    if not best or (slate and best[0] < max(2, len(slate) // 2)):
        print("export_fight_order: no UFC Stats event matches DKSalaries.csv; using the workbook.")
        return None
    print(f"export_fight_order: order from UFC Stats, {best[1]} ({best[0]} of {len(slate)} slate fighters matched)")
    return [(i + 1, b) for i, b in enumerate(best[2])], f"UFC Stats: {best[1]}"


def from_workbook(src):
    try:
        import openpyxl
    except ImportError:
        print("export_fight_order: openpyxl is not installed; skipping the workbook.")
        return None
    if not src.exists():
        print(f"export_fight_order: {src.name} not found.")
        return None
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True, keep_vba=False)
    if SHEET not in wb.sheetnames:
        print(f"export_fight_order: sheet '{SHEET}' not found in {src.name}.")
        return None
    fights = {}
    for i, row in enumerate(wb[SHEET].iter_rows(values_only=True)):
        if i == 0 or not row:
            continue
        name = str(row[0]).strip() if row[0] is not None else ""
        try:
            fid = int(float(row[1])) if len(row) > 1 and row[1] is not None else None
        except (TypeError, ValueError):
            fid = None
        if name and fid is not None and name not in fights.setdefault(fid, []):
            fights[fid].append(name)
    wb.close()
    fights = {k: v for k, v in fights.items() if v}
    if not fights:
        print("export_fight_order: no fighters with fight IDs in the workbook.")
        return None
    print(f"export_fight_order: order from {src.name} FightIDs")
    return [(fid, fights[fid]) for fid in sorted(fights)], src.name


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = Path(args[0]) if args else HERE / "Post Contest Sim.xlsm"
    if not src.is_absolute():
        src = Path.cwd() / src
    result = None
    if "--workbook" not in sys.argv:
        result = from_ufcstats(slate_names())
    if result is None:
        result = from_workbook(src)
    if result is None:
        print("export_fight_order: nothing written; the site keeps its start-time order.")
        return 0
    fights, source = result
    out = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "source": source,
        "fights": [{"fight_id": fid, "fighters": names} for fid, names in fights],
    }
    with open(HERE / "fight_order.json", "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    print(f"export_fight_order: wrote fight_order.json ({len(fights)} fights)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
